#!/usr/bin/env python3
"""ABS 成熟度条例索引。只写入本次打开过的条文。"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import json

OUT = "/Users/yalizhu/Documents/Dev/market-research/ABS成熟度条例索引.xlsx"
JSON_OUT = "/Users/yalizhu/Documents/Dev/market-research/web/src/data/cashloanAbsOrdinance.json"

CLAUSE_ID = {
    "法律能否出表": "legal",
    "发行是否连续": "issuance",
    "哪类资产进得了池": "asset",
    "买方是否稳定": "buyers",
    "发起人还要留多少": "retention",
    "发行基础设施": "plumbing",
    "监管是否允许": "policy",
}

INK = "232946"
CREAM = "F7F4EE"
GAP = "F3F0E8"
WHITE = "FFFFFF"
LINE = "D9D3C7"
LINK = "1F4E79"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
wrap = Alignment(wrap_text=True, vertical="top")
header_font = Font(name="Calibri", bold=True, color=WHITE, size=11)
body = Font(name="Calibri", size=11, color="222222")
gap_font = Font(name="Calibri", size=11, color="595959")
title_font = Font(name="Calibri", bold=True, size=16, color=INK)
note_font = Font(name="Calibri", size=11, color="333333")

HEADERS = [
    "国家",
    "代码",
    "载体",
    "条款",
    "镜头",
    "条例结论",
    "出处",
    "条号",
    "文本时点",
    "链接",
    "缺口",
]

# status: cited | gap
ROWS = []


def add(**kw):
    ROWS.append(kw)


# --- 印度 ---
IN = dict(
    country="印度",
    code="IN",
    vehicle="标准资产证券化（银行、小型银行、全印金融机构、含住房金融公司在内的非银行金融公司）",
    src="印度储备银行《标准资产证券化指引》RBI/DOR/2021-22/85",
    when="2021-09-24 发布，页面标注更新至 2022-12-05",
    url="https://www.rbi.org.in/scripts/bs_viewmasdirections.aspx?id=12165",
)

add(**IN, clause="法律能否出表", lens="本地", status="cited",
    finding="信托型特殊目的载体应破产隔离、无自由裁量。资本充足上的出表，还要求转让后的资产在发起人破产（点名《破产法》）时也离开发起人及其债权人，证券不是发起人的债务。",
    article="第 5 条(a) 款；第 30 条(e)(iii)；第 81 条(c)(d)")
add(**IN, clause="法律能否出表", lens="国际", status="gap",
    finding="印度银行海外分行做证券化，不得违反《外汇管理法》及其细则。指引没有写本地贷款能否卖给境外载体，也没有写本息能否汇出。",
    article="第 7 条",
    gap="不把第 7 条写成资本项目已放行。")
add(**IN, clause="发行是否连续", lens="本地", status="gap",
    finding="指引不发布消费贷、现金贷或信用卡证券的金额和单数。",
    article="—",
    gap="评级公司新闻稿不是条例，不写入本表。")
add(**IN, clause="发行是否连续", lens="国际", status="gap",
    finding="指引没有该国资产的美元或离岸证券发行记录。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**IN, clause="哪类资产进得了池", lens="本地", status="cited",
    finding="再证券化、合成证券化、循环额度（条文点名信用卡应收和现金信用）、指定期内的重组贷款、对其他贷款机构的敞口、全印金融机构的转贷款、本息都到期一次偿还的贷款、剩余期限短于 365 天的贷款，不能做底层。除此之外，表内且分类为正常的贷款和垫款可以做底层。",
    article="第 6 条(d)；第 8 条")
add(**IN, clause="哪类资产进得了池", lens="国际", status="gap",
    finding="指引没有写国际买方实际买过哪一类资产。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**IN, clause="买方是否稳定", lens="本地", status="cited",
    finding="单一投资者的最低票面是 1000 万卢比。向 50 人及以上发售，按指引所引的印度证监会 2008 年《证券化债务工具与证券收据发行及上市条例》必须上市。指引没有写银行、货币基金、保险或养老金是否反复配置。",
    article="第 28 条；第 29 条",
    gap="本次没有另开证监会 2008 年条例全文，只引用储备银行指引里点名的那一句。")
add(**IN, clause="买方是否稳定", lens="国际", status="gap",
    finding="指引没有写买方是本地账户还是国际账户，也没有相对国债的利差。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**IN, clause="发起人还要留多少", lens="本地", status="cited",
    finding="住房抵押证券以外：原始期限 24 个月及以内，最低自持为被证券化贷款账面价值的 5%；超过 24 个月，以及第 6 条但书允许的一次还本付息贷款，为 10%。住房抵押证券为 5%。发起人对同一结构的证券化敞口合计不超过该结构证券化敞口的 20%，利息剥离不计入这 20%。这是监管下限和上限，不是成交里实际留下的比例。",
    article="第 12 条；第 13 条；第 25 条")
add(**IN, clause="发起人还要留多少", lens="国际", status="gap",
    finding="指引没有写国际发行里发起人或本地银行要担保或认购多少。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**IN, clause="发行基础设施", lens="本地", status="cited",
    finding="信托文件要写明受托人的职能、权利义务和投资者权利，投资者可随时更换受托人。指引没有规定披露是池级还是逐笔，也没有规定滞后天数。",
    article="第 30 条(e)(iv)",
    gap="不补披露频率。")
add(**IN, clause="发行基础设施", lens="国际", status="gap",
    finding="指引没有要求使用国际评级，也没有写跨境买方的报告频率。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**IN, clause="监管是否允许", lens="本地", status="cited",
    finding="适用主体是表列商业银行（不含地区农村银行）、全印定期金融机构、小型金融银行，以及含住房金融公司在内的全部非银行金融公司。正常类表内贷款可以证券化；第 6 条列出的资产不可以。",
    article="第 3 条；第 6 条；第 8 条")
add(**IN, clause="监管是否允许", lens="国际", status="gap",
    finding="第 7 条只要求海外分行的证券化不得违反《外汇管理法》。没有写外资能否持有这些证券，也没有写本息能否汇出。",
    article="第 7 条",
    gap="不把外汇管理法的转引写成汇出已放行。")

# --- 印尼 KIK-EBA ---
KIK = dict(
    country="印度尼西亚",
    code="ID",
    vehicle="集合投资合同型资产支持证券（KIK-EBA）",
    src="金融服务管理局条例 65/POJK.04/2017",
    when="2017 年条例。条号来自法规汇编站对该条例的收录，管理局目录页标题一致。没有打开 PDF 原件。",
    url="https://ojk.go.id/id/regulasi/Pages/Pedoman-Penerbitan-dan-Pelaporan-Efek-Beragun-Aset-Berbentuk-Kontrak-Investasi-Kolektif.aspx",
)
add(**KIK, clause="法律能否出表", lens="本地", status="cited",
    finding="资产须以法律上的断卖或断换转入集合投资合同。合同要努力确认资产与发起人分离、发起人破产时不进入破产财产；发起人让出全部权利、不再持有、不控制该合同；合同对资产损失没有向发起人追索的权利。法律上的断卖须有在管理局注册的法律顾问意见。会计上的断卖不是强制，若做了必须一贯并有注册会计意见。",
    article="第 2 条第(4)款至第(8)款")
add(**KIK, clause="法律能否出表", lens="国际", status="gap",
    finding="条例没有写本地贷款能否卖给境外载体，也没有写资本项目是否放行。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**KIK, clause="发行是否连续", lens="本地", status="gap",
    finding="条例不发布近 1–3 年的金额和单数。",
    article="—",
    gap="存管月报里的资产支持证券行不区分消费贷，不写入本表。")
add(**KIK, clause="发行是否连续", lens="国际", status="gap",
    finding="条例没有离岸或美元证券的发行记录。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**KIK, clause="哪类资产进得了池", lens="本地", status="cited",
    finding="组合可以包括商业票据产生的应收、信用卡应收、未来应收、发放信贷产生的应收、政府担保的债务证券、增信手段、未来现金流及其权利、未来收入及其权利，以及其他相当的金融资产。资产要能产生现金流，发起人在法律上拥有或控制，并且可以自由转让。条例没有把无抵押现金贷单列出来，也没有把它排除。",
    article="第 2 条第(1)款至第(3)款")
add(**KIK, clause="哪类资产进得了池", lens="国际", status="gap",
    finding="条例没有写国际买方实际买过哪一类。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**KIK, clause="买方是否稳定", lens="本地", status="gap",
    finding="招募说明书要说明证券适合哪类机构投资者。条例没有写银行、货币基金、保险或养老金是否反复买入。",
    article="第 20 条",
    gap="不把「适合某类机构投资者」写成买方已经稳定。")
add(**KIK, clause="买方是否稳定", lens="国际", status="gap",
    finding="招募说明书要披露境内外投资者的税务处理。这是披露项目，不是买方名单，也没有利差。",
    article="第 20 条(h)",
    gap="不把税务披露写成国际账户已经在买。")
add(**KIK, clause="发起人还要留多少", lens="本地", status="cited",
    finding="发起人买回或断换其已经转让资产，最多为已转让金融资产价值的 10%。这是买回上限，不是最低自持比例。增信手段可以包括同一合同内不同档的次级安排，条例没有规定次级档必须占多少。",
    article="第 3 条；第 2 条第(2)款(a)")
add(**KIK, clause="发起人还要留多少", lens="国际", status="gap",
    finding="条例没有写国际发行中的担保或认购比例。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**KIK, clause="发行基础设施", lens="本地", status="cited",
    finding="合同由在管理局注册的公证人做成公证文书，并写明管理人、托管银行、服务商、会计师、法律顾问；公开发行时还要写明评级机构。资产以托管银行名义、为证券持有人登记。持有人报告最迟于次月 12 日提交，其中包含分档资产报告和按市场利率、最近评级、预期付款估计的公允价值。年报最迟于次年 3 月 31 日提交。向管理局报送月报和年报。条例没有要求逐笔披露。",
    article="第 4 条；第 5 条；第 6 条；第 26 条；第 27 条；第 29 条")
add(**KIK, clause="发行基础设施", lens="国际", status="gap",
    finding="评级机构是获得管理局许可的公司。条例没有区分本地评级和国际评级，也没有单写跨境买方的报告频率。",
    article="第 20 条(i)",
    gap="不把管理局许可的评级机构写成国际评级。")
add(**KIK, clause="监管是否允许", lens="本地", status="cited",
    finding="信贷产生的应收和信用卡应收都在第 2 条的许可范围内。条例没有另写一条禁止非银或现金贷入池。",
    article="第 2 条第(1)款(b)(d)",
    gap="许可范围不是已经发行的证明。")
add(**KIK, clause="监管是否允许", lens="国际", status="gap",
    finding="条例没有写外资能否持有这些证券，也没有写本息能否汇出。",
    article="—",
    gap="已打开的条文没有这一读数。")

# --- 印尼 EBA-SP ---
SP = dict(
    country="印度尼西亚",
    code="ID",
    vehicle="参与凭证型资产支持证券（EBA-SP，住房二次融资）",
    src="金融服务管理局条例 23/POJK.04/2014，第 9 条由 20/POJK.04/2017 修改",
    when="20/POJK.04/2017 的官方 PDF 已打开。23/POJK.04/2014 的资产定义读自法规汇编站收录。2026-09-17 有一份编号 12 的住房参与凭证新条例页面，第三方标注尚未审完，本次不改挂到那份条例。",
    url="https://www.ojk.go.id/id/kanal/pasar-modal/regulasi/peraturan-ojk/Documents/Pages/POJK-Nomor-20-POJK.04-2017/SAL%20POJK%2020.pdf",
)
add(**SP, clause="法律能否出表", lens="本地", status="cited",
    finding="形成组合的金融资产，须由发行人以法律上的断卖从原债权人取得，再断卖给持有人；或者由发行人为持有人的利益向原债权人断卖取得。断卖须有法律顾问意见，并符合通用会计准则、一贯执行、有会计意见。",
    article="20/POJK.04/2017 修改后的第 9 条第(1)款至第(3)款")
add(**SP, clause="法律能否出表", lens="国际", status="gap",
    finding="已打开的第 9 条没有写卖给境外载体或资本项目。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="发行是否连续", lens="本地", status="gap",
    finding="条例不发布金额和单数。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="发行是否连续", lens="国际", status="gap",
    finding="条例没有离岸发行记录。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="哪类资产进得了池", lens="本地", status="cited",
    finding="金融资产是原债权人发放住房按揭贷款所产生的应收，包括附着的担保和抵押权。这不是现金贷，也不是信用卡。",
    article="23/POJK.04/2014 第 1 条",
    gap="资产定义来自 23/2014 的汇编收录，不是 20/2017 那份已打开的 PDF。")
add(**SP, clause="哪类资产进得了池", lens="国际", status="gap",
    finding="已打开的条文没有国际买方买过的资产类别。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="买方是否稳定", lens="本地", status="gap",
    finding="已打开的第 9 条没有买方类型。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="买方是否稳定", lens="国际", status="gap",
    finding="已打开的第 9 条没有国际买方或利差。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="发起人还要留多少", lens="本地", status="cited",
    finding="原债权人买入参与凭证，最多为组合总价值的 10%。第 9 条第(1)款(a) 路径下的发行人，在首次发行时买入也最多 10%。第(1)款(b) 路径下，若首次发行没有被市场全部吸收，发行人买入可以超过 10%。这是买入上限，不是最低自持。",
    article="20/POJK.04/2017 修改后的第 9 条第(4)款、(4a)款、(4b)款")
add(**SP, clause="发起人还要留多少", lens="国际", status="gap",
    finding="已打开的第 9 条没有国际发行的担保或认购比例。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="发行基础设施", lens="本地", status="gap",
    finding="已打开的第 9 条要求法律顾问意见和会计意见，没有写托管、登记、披露频率。",
    article="第 9 条第(2)款、第(3)款",
    gap="不从其他条例补基础设施。")
add(**SP, clause="发行基础设施", lens="国际", status="gap",
    finding="已打开的第 9 条没有国际评级或跨境报告频率。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**SP, clause="监管是否允许", lens="本地", status="cited",
    finding="这条轨道的资产是住房按揭应收。它不授权现金贷入池。",
    article="23/POJK.04/2014 第 1 条")
add(**SP, clause="监管是否允许", lens="国际", status="gap",
    finding="已打开的条文没有外资持有或汇出。",
    article="—",
    gap="已打开的条文没有这一读数。")

# --- 菲律宾 ---
PH = dict(
    country="菲律宾",
    code="PH",
    vehicle="特别目的载体发行的资产支持证券",
    src="共和国法 9267 号《2004 年证券化法》",
    when="2004-03-19。文本读自最高法院电子图书馆。",
    url="https://elibrary.judiciary.gov.ph/thebookshelf/showdocs/2/1499",
)
add(**PH, clause="法律能否出表", lens="本地", status="cited",
    finding="证券化定义为卖方无追索地把资产卖给特别目的载体。真实出售要同时满足：资产与发起人或卖方及其债权人隔离；载体可以质押、抵押或交换；转让人放弃有效控制；转让是出售、让与或交换且无追索；载体享有收益和处分权；转让人不能取回资产，载体也不能要求退回对价；载体承担资产风险。陈述与保证不破坏最后一项。",
    article="第 3 条第 1 款；第 12 条")
add(**PH, clause="法律能否出表", lens="国际", status="gap",
    finding="已打开的条文没有写卖给境外载体，也没有写外汇或资本项目。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**PH, clause="发行是否连续", lens="本地", status="gap",
    finding="第 4 条要求证监会在年报中列出特别目的载体以及证券化资产的类型和金额。法条本身没有金额。本次没有找到这份年报。",
    article="第 4 条",
    gap="不借用新闻里的住房债券金额。")
add(**PH, clause="发行是否连续", lens="国际", status="gap",
    finding="法条没有离岸发行记录。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**PH, clause="哪类资产进得了池", lens="本地", status="cited",
    finding="资产是具有预期现金付款流的贷款、应收或其他类似金融资产，包括应收、按揭贷款和其他债务工具。未来才产生的应收须经证监会或中央银行批准。排除中央或地方将来的特许权使用费、收费或税捐。法条没有点名无抵押现金贷。",
    article="第 3 条第 3 款")
add(**PH, clause="哪类资产进得了池", lens="国际", status="gap",
    finding="法条没有写国际买方实际买过哪一类。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**PH, clause="买方是否稳定", lens="本地", status="gap",
    finding="资产支持证券须按《证券监管法》第 8 条、第 12 条登记；符合第 9 条、第 10 条的，改为提交通知和披露声明。法条没有买方类型。",
    article="第 7 条",
    gap="已打开的条文没有这一读数。")
add(**PH, clause="买方是否稳定", lens="国际", status="gap",
    finding="法条没有国际买方或利差。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**PH, clause="发起人还要留多少", lens="本地", status="gap",
    finding="第 12 条第 6 款禁止转让人取回资产。这是真实出售的条件，不是自持比例。已打开的条文没有次级档或自持百分比。",
    article="第 12 条第 6 款",
    gap="不把禁止取回写成自持比例。")
add(**PH, clause="发起人还要留多少", lens="国际", status="gap",
    finding="法条没有国际发行的担保或认购比例。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**PH, clause="发行基础设施", lens="本地", status="cited",
    finding="资产支持证券发行前须由经认证的信用评级机构评级。评级机构开始评级前须取得证监会认证。证监会至少每三年检查一次评级机构。法条没有规定池级或逐笔，也没有滞后天数。",
    article="第 43 条；第 45 条；第 46 条",
    gap="不补披露频率。")
add(**PH, clause="发行基础设施", lens="国际", status="gap",
    finding="评级机构由菲律宾证监会认证。法条没有要求国际评级，也没有写跨境报告频率。",
    article="第 45 条",
    gap="不把本地认证写成国际评级。")
add(**PH, clause="监管是否允许", lens="本地", status="cited",
    finding="贷款和应收可以成为资产池。若发起人是银行或其他受中央银行监管的金融中介，或其关联方，或者载体是特别目的信托，还需要中央银行事先背书。法条没有单独禁止现金贷入池。",
    article="第 3 条第 3 款；第 11 条")
add(**PH, clause="监管是否允许", lens="国际", status="gap",
    finding="法条没有写外资能否持有这些证券，也没有写本息能否汇出。",
    article="—",
    gap="已打开的条文没有这一读数。")

# --- 泰国 ---
TH = dict(
    country="泰国",
    code="TH",
    vehicle="特别目的法人或信托",
    src="佛历 2540 年特别目的法人证券化皇家法令，经佛历 2558 年修订法替换所列条文",
    when="修订法 2015-03-20 制定，政府公报第 132 卷 21 Kor 分册 2015-03-26 公布，公布次日生效。本次打开的是这份修订法的公报文本，其中替换了原法第 3、5、9 至 24、29、30、34 条。没有打开 1997 年未被替换的其余条文，也没有打开证监会的资产类别公告。",
    url="https://web.senate.go.th/bill/bk_data/60-6.PDF",
)
add(**TH, clause="法律能否出表", lens="本地", status="cited",
    finding="证券化是特别目的法人从资产出售人受让资产，或接受资产作为担保，向投资者发行证券，并用所得向出售人付款；对证券持有人的偿付取决于所受让或所担保资产的收入。已打开的条文没有真实出售或破产隔离的判定标准。第 15 条只规定两类债权转让可以不通知债务人即生效。第 20 条是转让价格被视为公平、不构成低价损害债权人的两种估价。",
    article="修订后第 3 条；第 15 条；第 20 条",
    gap="不把第 15 条或第 20 条写成真实出售已经成立。")
add(**TH, clause="法律能否出表", lens="国际", status="gap",
    finding="已打开的条文没有写卖给境外载体或资本项目。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="发行是否连续", lens="本地", status="gap",
    finding="已打开的条文没有金额和单数。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="发行是否连续", lens="国际", status="gap",
    finding="已打开的条文没有离岸发行记录。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="哪类资产进得了池", lens="本地", status="cited",
    finding="资产是能够产生收入的金钱债权，或者按照证监会公告标准、能够产生收入的未来金钱债权。证监会有权规定允许证券化的资产类型和据此发行的证券类型。已打开的法律文本没有点名现金贷、信用卡或车贷。",
    article="修订后第 3 条；第 5 条第(3)款",
    gap="资产白名单在证监会公告里。本次没有打开那些公告，不写某一类贷款已经获准。")
add(**TH, clause="哪类资产进得了池", lens="国际", status="gap",
    finding="已打开的条文没有国际买方实际买过的类别。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="买方是否稳定", lens="本地", status="gap",
    finding="已打开的条文没有买方类型。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="买方是否稳定", lens="国际", status="gap",
    finding="已打开的条文没有国际买方或利差。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="发起人还要留多少", lens="本地", status="gap",
    finding="特别目的法人清偿证券持有人之后，把剩余资产返还出售人或解除担保，法人即终止。这是剩余资产的归属，不是自持比例。已打开的条文没有次级档或自持百分比。",
    article="修订后第 24 条第(1)款",
    gap="不把剩余资产返还写成自持比例。")
add(**TH, clause="发起人还要留多少", lens="国际", status="gap",
    finding="已打开的条文没有国际发行的担保或认购比例。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="发行基础设施", lens="本地", status="cited",
    finding="特别目的法人可以是有限公司、公众公司、证监会公告指定的其他法人，或依资本市场信托法设立的信托，且只能为证券化而设。项目连同新证券发行许可申请一并提交证监会办公室。收款代理人须为已转让资产单设账户和债务人名单，债务人可以查阅自己的信息。已打开的条文没有评级要求，也没有披露滞后天数。",
    article="修订后第 9 条；第 10 条；第 15/1 条",
    gap="不补评级或披露频率。")
add(**TH, clause="发行基础设施", lens="国际", status="gap",
    finding="已打开的条文没有国际评级或跨境报告频率。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**TH, clause="监管是否允许", lens="本地", status="cited",
    finding="允许证券化的资产类型由证监会规定，不在本次打开的法律文本里。若特别目的法人的活动在性质上属于金融业务或信贷地产金融业务，它可以从事该活动而无需再取得那些法律下的许可。这是载体的许可豁免，不是现金贷发起人的许可。",
    article="修订后第 5 条第(3)款；第 14 条")
add(**TH, clause="监管是否允许", lens="国际", status="gap",
    finding="已打开的条文没有外资持有或汇出。",
    article="—",
    gap="已打开的条文没有这一读数。")

# --- 墨西哥 ---
MX = dict(
    country="墨西哥",
    code="MX",
    vehicle="不可撤销信托发行的信托型上市凭证",
    src="证券市场法",
    when="2005-12-30 刊登于联邦官方日报。众议院文本标注最近一次改革刊登于 2025-11-14。本次打开的是众议院现行文本，只读了上市凭证专章。没有打开《信用工具与操作总法》的信托章。",
    url="https://www.diputados.gob.mx/LeyesBiblio/pdf/LMV.pdf",
)
add(**MX, clause="法律能否出表", lens="本地", status="cited",
    finding="上市凭证可以通过不可撤销信托发行，名称应为信托型上市凭证。同一信托下不同系列的账户可以约定只用于本系列，即使发行信托进入商业破产或破产，也不能用于其他系列。这是系列之间的隔离，不是发起人资产的真实出售或破产隔离。已打开的专章没有真实出售标准。",
    article="第 63 条；第 64 条",
    gap="不把第 64 条的系列隔离写成发起人出表。")
add(**MX, clause="法律能否出表", lens="国际", status="gap",
    finding="第 61 条允许有权签署信用工具的外国法人发行上市凭证。这是外国发行人，不是本地贷款卖给境外载体，也不是资本项目放行。",
    article="第 61 条",
    gap="不把外国发行人条款写成跨境出表。")
add(**MX, clause="发行是否连续", lens="本地", status="gap",
    finding="已打开的专章没有金额和单数。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**MX, clause="发行是否连续", lens="国际", status="gap",
    finding="已打开的专章没有离岸发行记录。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**MX, clause="哪类资产进得了池", lens="本地", status="cited",
    finding="上市凭证代表对集体信贷的个别参与，或对信托财产的第 63 条权利：信托财产的所有权或名义份额、果实和剩余价值、出售所得，以及本金、利息或其他款项。凭证可以是优先或次级，持有人之间的受偿顺序可以不同。专章没有点名消费贷、现金贷或信用卡。第 63 Bis 1 条的开发、不动产、指数型凭证是另一类工具，不计入贷款证券化。",
    article="第 62 条；第 63 条")
add(**MX, clause="哪类资产进得了池", lens="国际", status="gap",
    finding="已打开的专章没有国际买方实际买过的类别。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**MX, clause="买方是否稳定", lens="本地", status="gap",
    finding="发行信托型上市凭证的受托人只能是信贷机构、经纪商，以及投资基金的运营公司。这是受托人资格，不是买方构成。",
    article="第 63 条末段",
    gap="不把受托人资格写成买方稳定。")
add(**MX, clause="买方是否稳定", lens="国际", status="gap",
    finding="已打开的专章没有国际买方或利差。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**MX, clause="发起人还要留多少", lens="本地", status="gap",
    finding="第 62 条允许优先档和次级档。已打开的专章没有次级档比例或发起人自持百分比。",
    article="第 62 条",
    gap="不把可以分档写成已经规定自持比例。")
add(**MX, clause="发起人还要留多少", lens="国际", status="gap",
    finding="已打开的专章没有国际发行的担保或认购比例。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**MX, clause="发行基础设施", lens="本地", status="cited",
    finding="凭证须载明发行人、发行金额、持有人权利、利率或收益、期限、提前到期、担保，以及发起人、信托委托人和信托财产管理人的义务。凭证存放于本法监管的证券存管机构，可以附息票。已打开的专章没有评级要求，也没有披露滞后天数。",
    article="第 64 条",
    gap="不补评级或披露频率。")
add(**MX, clause="发行基础设施", lens="国际", status="gap",
    finding="已打开的专章没有国际评级或跨境报告频率。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**MX, clause="监管是否允许", lens="本地", status="gap",
    finding="已打开的专章没有写监管允许或限制消费信贷出表，也没有写非银或现金贷能否入池。",
    article="—",
    gap="已打开的条文没有这一读数。")
add(**MX, clause="监管是否允许", lens="国际", status="gap",
    finding="已打开的专章没有写外资能否持有这些证券，也没有写本息能否汇出。",
    article="—",
    gap="已打开的条文没有这一读数。")


def style_header(ws, row, ncol):
    fill = PatternFill("solid", fgColor=INK)
    for col in range(1, ncol + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    ws.row_dimensions[row].height = 22
    ws.auto_filter.ref = f"A{row}:{get_column_letter(ncol)}{row}"
    ws.freeze_panes = f"A{row + 1}"
    ws.auto_filter.ref = None  # set after data
    ws.sheet_view.showGridLines = False


def write_table(ws, rows, start_row):
    for c, h in enumerate(HEADERS, 1):
        ws.cell(start_row, c, h)
    style_header(ws, start_row, len(HEADERS))
    for i, r in enumerate(rows):
        excel_row = start_row + 1 + i
        values = [
            r["country"],
            r["code"],
            r["vehicle"],
            r["clause"],
            r["lens"],
            r["finding"],
            r["src"],
            r["article"],
            r["when"],
            r["url"],
            r.get("gap") or "",
        ]
        is_gap = r["status"] == "gap"
        fill = PatternFill("solid", fgColor=GAP if is_gap else WHITE)
        font = gap_font if is_gap else body
        for c, val in enumerate(values, 1):
            cell = ws.cell(excel_row, c, val)
            cell.font = font
            cell.fill = fill
            cell.alignment = wrap
            cell.border = thin
            if c == 10 and val:
                cell.hyperlink = val
                cell.font = Font(name="Calibri", size=11, color=LINK, underline="single")
        ws.row_dimensions[excel_row].height = 78
    last = start_row + len(rows)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(HEADERS))}{last}"
    ws.freeze_panes = f"A{start_row + 1}"
    widths = [14, 8, 36, 18, 10, 54, 42, 36, 36, 28, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "ABS 成熟度条例索引"
    ws.print_title_rows = f"{start_row}:{start_row}"
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.zoomScale = 110


def main():
    wb = Workbook()
    cover = wb.active
    cover.title = "说明"
    cover.sheet_view.showGridLines = False
    cover.sheet_view.zoomScale = 120
    cover.column_dimensions["A"].width = 28
    cover.column_dimensions["B"].width = 110
    lines = [
        ("ABS 成熟度条例索引", ""),
        ("生成日期", "2026-10-07"),
        ("索引", "国家。只收已投五国：印度、印度尼西亚、菲律宾、泰国、墨西哥。"),
        ("分类", "七个条款，每个条款分本地与国际。工作表「按条款」把同一批行按条款排在一起。"),
        ("写入原则", "只写本次打开过的法律条文。条例没有写到的，留在「缺口」，不补金额、单数、买方占比或利差。"),
        ("不写入的", "印度 CRISIL 的证券化成交额和资产占比是评级公司新闻稿，不是条例。印尼存管月报的资产支持证券行不区分消费贷。新闻里的单笔住房债券不写入。"),
        ("印度文本", "储备银行指引页面，RBI/DOR/2021-22/85，2021-09-24，页面标注更新至 2022-12-05。"),
        ("印尼文本", "KIK-EBA 为 65/POJK.04/2017，管理局目录页对得上标题，条号来自法规汇编站收录，PDF 原件本次没打开。EBA-SP 的第 9 条来自 20/POJK.04/2017 官方 PDF；住房按揭的定义来自 23/POJK.04/2014 的汇编收录。"),
        ("菲律宾文本", "共和国法 9267 号，2004-03-19，最高法院电子图书馆。"),
        ("泰国文本", "2015-03-26 政府公报上的佛历 2558 年修订法，替换了 1997 年皇家法令的所列条文。1997 年其余条文和证监会资产公告本次没打开。"),
        ("墨西哥文本", "众议院《证券市场法》现行文本，最近改革刊登日 2025-11-14。只读了上市凭证专章。《信用工具与操作总法》的信托章本次没打开。"),
        ("怎么读缺口", "灰色行表示这条条例没有给出该条款要的读数。白色行是条文里写明的内容，仍可能在「缺口」里标明不能外推的边界。"),
    ]
    cover["A1"] = lines[0][0]
    cover["A1"].font = title_font
    cover.merge_cells("A1:B1")
    cover.row_dimensions[1].height = 28
    for i, (k, v) in enumerate(lines[1:], 3):
        cover.cell(i, 1, k).font = Font(name="Calibri", bold=True, size=11, color=INK)
        cover.cell(i, 1).alignment = Alignment(vertical="top")
        cover.cell(i, 1).fill = PatternFill("solid", fgColor=CREAM)
        cell = cover.cell(i, 2, v)
        cell.font = note_font
        cell.alignment = wrap
        cover.row_dimensions[i].height = 48
    cover.row_dimensions[8].height = 62
    cover.page_setup.orientation = "landscape"
    cover.page_setup.fitToPage = True
    cover.page_setup.fitToWidth = 1
    cover.page_setup.fitToHeight = 1
    cover.sheet_properties.pageSetUpPr.fitToPage = True

    order = ["印度", "印度尼西亚", "菲律宾", "泰国", "墨西哥"]
    for name in order:
        ws = wb.create_sheet(name)
        subset = [r for r in ROWS if r["country"] == name]
        write_table(ws, subset, 1)

    clause_order = [
        "法律能否出表",
        "发行是否连续",
        "哪类资产进得了池",
        "买方是否稳定",
        "发起人还要留多少",
        "发行基础设施",
        "监管是否允许",
    ]
    lens_order = ["本地", "国际"]
    by_clause = sorted(
        ROWS,
        key=lambda r: (
            clause_order.index(r["clause"]),
            lens_order.index(r["lens"]),
            order.index(r["country"]),
            r["vehicle"],
        ),
    )
    ws = wb.create_sheet("按条款", 1)
    write_table(ws, by_clause, 1)

    wb.save(OUT)
    vehicles: dict[str, list] = {}
    current = None
    bucket = None
    for r in ROWS:
        key = (r["code"], r["vehicle"])
        if key != current:
            current = key
            bucket = {
                "name": r["vehicle"],
                "source": r["src"],
                "asOf": r["when"],
                "url": r["url"],
                "readings": [],
            }
            vehicles.setdefault(r["code"], []).append(bucket)
        bucket["readings"].append({
            "clause": CLAUSE_ID[r["clause"]],
            "lens": "local" if r["lens"] == "本地" else "intl",
            "finding": r["finding"],
            "article": r["article"],
            "gap": r.get("gap") or "",
            "cited": r["status"] == "cited",
        })
    payload = {
        "reviewed": "2026-10-07",
        "scope": "已投五国的证券化专法。发行金额、买方占比和利差不是条例，不在此列。",
        "byCode": vehicles,
    }
    with open(JSON_OUT, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(OUT, len(ROWS), JSON_OUT)


if __name__ == "__main__":
    main()
